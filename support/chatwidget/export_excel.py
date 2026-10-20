"""
export_excel.py — converts the three Kyambogo programme Excel files into one
JSON file (programmes_export.json) that import_excel.php then chunks,
embeds and loads into programme_chunks.

Mapping per file type:
  undergraduate.xlsx          -> ug-<code>  (overview + admission_requirements)
  graduate programmes.xlsx    -> pg-<code>  (overview incl. fees + admission_requirements)
  affiliated programmes.xlsx  -> aff-<no>   (one overview chunk per institution)

Output shape matches what sync_publish.php's chunk_programme() expects:
  { "id", "name", "overview", "structure", "admission_requirements", "career_prospects" }
"""

import json
import openpyxl
import re

BASE = r"C:\Users\HP\Downloads"
OUT = r"C:\Users\HP\Downloads\chatwidget\programmes_export.json"


def clean(v):
    """None/empty -> '', otherwise stripped text."""
    if v is None:
        return ""
    return str(v).strip()


def slugify(text):
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s[:40] or "x"


def parse_undergraduate(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or not any(row):
            continue
        name = clean(row[1])
        code_day = clean(row[2])
        code_eve = clean(row[3])
        reqs = clean(row[4])
        duration = clean(row[5])
        school = clean(row[6])
        if not name:
            continue
        overview_parts = [f"{name} is a {duration} programme" if duration else name]
        overview_parts.append(f"offered by the {school}" if school else "offered at Kyambogo University")
        codes = []
        if code_day:
            codes.append(f"day programme code {code_day}")
        if code_eve:
            codes.append(f"evening programme code {code_eve}")
        if codes:
            overview_parts.append("(" + ", ".join(codes) + ")")
        overview = " ".join(overview_parts) + "."
        out.append({
            "id": f"ug-{slugify(code_day or name)}",
            "name": name,
            "overview": overview,
            "structure": "",
            "admission_requirements": reqs,
            "career_prospects": "",
        })
    wb.close()
    return out


def parse_graduate(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or not any(row):
            continue
        name = clean(row[1])
        code = clean(row[2])
        fees = clean(row[3])
        intl = clean(row[4])
        reqs = clean(row[5])
        faculty = clean(row[6])
        if not name:
            continue
        parts = [f"{name} is offered by the {faculty}" if faculty else name]
        fee_parts = []
        if fees:
            fee_parts.append(f"fees of {fees} for national and East African students")
        if intl:
            fee_parts.append(f"{intl} for international students")
        if fee_parts:
            parts.append("with " + " and ".join(fee_parts))
        overview = " ".join(parts) + "."
        out.append({
            "id": f"pg-{slugify(code or name)}",
            "name": name,
            "overview": overview,
            "structure": "",
            "admission_requirements": reqs,
            "career_prospects": "",
        })
    wb.close()
    return out


def parse_affiliated(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or not any(row):
            continue
        institution = clean(row[1])
        ownership = clean(row[2])
        location = clean(row[3])
        programmes = clean(row[4])
        status = clean(row[5])
        if not institution:
            continue
        parts = [f"{institution} is a {ownership} institution" if ownership else institution]
        if location:
            parts.append(f"located in {location}")
        parts.append("affiliated to Kyambogo University")
        if status:
            parts.append(f"with affiliation status: {status}")
        text = " ".join(parts) + "."
        if programmes:
            text += f" It offers: {programmes}."
        out.append({
            "id": f"aff-{i}-{slugify(institution)}",
            "name": institution,
            "overview": text,
            "structure": "",
            "admission_requirements": "",
            "career_prospects": "",
        })
    wb.close()
    return out


def main():
    programmes = []
    programmes += parse_undergraduate(BASE + r"\undergraduate.xlsx")
    programmes += parse_graduate(BASE + r"\graduate programmes.xlsx")
    programmes += parse_affiliated(BASE + r"\affiliated programmes.xlsx")

    # ids must stay unique - if two programmes slugified to the same id,
    # disambiguate with a numeric suffix
    seen = {}
    for p in programmes:
        base = p["id"]
        n = seen.get(base, 0)
        if n:
            p["id"] = f"{base}-{n + 1}"
        seen[base] = n + 1

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(programmes, f, ensure_ascii=False, indent=2)

    by_type = {}
    for p in programmes:
        t = p["id"].split("-")[0]
        by_type[t] = by_type.get(t, 0) + 1
    print(f"Exported {len(programmes)} programmes -> {OUT}")
    print("by type:", by_type)
    for p in programmes[:2]:
        print("sample:", json.dumps(p, ensure_ascii=False)[:300])


if __name__ == "__main__":
    main()
